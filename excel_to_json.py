"""F126.xlsx -> data.json çevirici.
Kullanım: python excel_to_json.py
Çıktı: ./data.json (site bu dosyayı okur)
F1 API'ye geçince bu script'e gerek kalmaz, providers/f1ApiProvider.js devreye girer.
"""
import json
import datetime
from pathlib import Path
import openpyxl

SRC = Path(r"C:\Users\Eren Aygen\Downloads\F126.xlsx")
OUT = Path(__file__).parent / "data.json"

def parse_value(v):
    if v is None or v == "":
        return None
    s = str(v).strip()
    if s in ("DNS", "DNF", "DSQ", "Ret"):
        return s
    try:
        # "10.0" -> 10, "7.5" -> 7.5
        f = float(s.replace(",", "."))
        return int(f) if f.is_integer() else f
    except ValueError:
        return s

def num(v):
    return v if isinstance(v, (int, float)) else 0

def dedup(headers):
    seen = {}
    out = []
    for h in headers:
        if h in seen:
            seen[h] += 1
            out.append(f"{h}_{seen[h]}")
        else:
            seen[h] = 1
            out.append(h)
    return out

def main():
    wb = openpyxl.load_workbook(SRC, read_only=True, data_only=True)

    # --- Teams ---
    ws = wb["Teams"]
    rows = list(ws.iter_rows(values_only=True))
    team_header = [str(c).strip() for c in rows[0][1:]]  # AUS..UAE, Total
    # "Total" kolonunu ayır
    team_races = [c for c in team_header if c.lower() != "total"]
    teams = []
    for r in rows[1:]:
        name = r[0] if len(r) > 0 else None
        # ilk boş satırda dur: Excel'de alt bloklar (motor, kümülatif) var
        if (name is None or str(name).strip() == "") and (len(r) < 2 or r[1] is None or str(r[1]).strip() == ""):
            break
        if not name or str(name).strip().lower() in ("team", "motor ureticileri"):
            continue
        vals = [parse_value(v) for v in list(r)[1:1 + len(team_header)]]
        results = dict(zip(team_header, vals))
        total = sum(num(v) for k, v in results.items() if k.lower() != "total")
        # Excel'de Total kolonu varsa ve sayısal ise onu tercih et
        xl_total = results.get("Total")
        if isinstance(xl_total, (int, float)) and xl_total > 0:
            total = xl_total
        teams.append({
            "name": str(name).strip(),
            "results": {k: v for k, v in results.items() if k.lower() != "total"},
            "total": total,
        })
    teams.sort(key=lambda t: t["total"], reverse=True)

    # --- Drivers ---
    ws = wb["Drivers"]
    rows = list(ws.iter_rows(values_only=True))
    raw_header = [(str(c).strip() if c is not None else "") for c in rows[0]]
    # ilk 2 kolon: No, Driver
    race_header = dedup(raw_header[2:])
    # "Total" kolonunu ayır
    race_names = [c for c in race_header if c.lower() != "total" and c != ""]
    drivers = []
    for r in rows[1:]:
        r = list(r) + [None] * (2 + len(race_header) - len(list(r)))
        # ilk boş satırda dur (alt blok kümülatif tablo)
        if (r[1] is None or str(r[1]).strip() == ""):
            break
        if str(r[1]).strip().lower() in ("driver", "sürücü", "s�r�c�"):
            continue
        number = parse_value(r[0])
        name = str(r[1]).strip()
        vals = [parse_value(v) for v in list(r[2:2 + len(race_header)])]
        # eksik kolonları doldur
        while len(vals) < len(race_header):
            vals.append(None)
        results = dict(zip(race_header, vals))
        xl_total = results.pop("Total", None)
        # boş isimli kolonları at
        results = {k: v for k, v in results.items() if k}
        total = sum(num(v) for v in results.values())
        if isinstance(xl_total, (int, float)) and xl_total > 0:
            total = xl_total
        drivers.append({
            "number": number,
            "name": name,
            "results": results,
            "total": total,
        })
    drivers.sort(key=lambda d: d["total"], reverse=True)

    data = {
        "source": "excel:F126.xlsx",
        "generatedAt": datetime.datetime.now().isoformat(timespec="seconds"),
        "races": race_names,
        "teamRaces": team_races,
        "drivers": drivers,
        "teams": teams,
    }
    OUT.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"OK: {len(drivers)} pilot, {len(teams)} takım, {len(race_names)} yarış -> {OUT}")

if __name__ == "__main__":
    main()
